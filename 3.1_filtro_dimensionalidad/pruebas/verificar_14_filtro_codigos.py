"""Prueba independiente del filtro de dimensionalidad (actividad 3.1, script 14).

Vuelve a calcular, sin usar ninguna función del script 14 y por otro camino, lo que el script dejó en
data/processed/filtro_codigos.parquet y en reportes/14_codigos_retenidos.csv:

  1. Versión y archivos. El conjunto de desarrollo es el registrado y los archivos del script tienen las
     columnas esperadas, sin identificadores de egresos ni de pacientes.
  2. Totales. Los casos de estudio, los casos con dato y los reingresos de cada variable, y los códigos
     distintos de cada bloque, contados directamente con pyarrow.
  3. Muestra de códigos. Para 300 códigos al azar, los que quedaron justo en el límite de los umbrales
     y cinco ejemplos conocidos, marca los casos que tienen el código revisando columna por columna
     (sin la tabla larga del script) y calcula la V como la correlación entre dos columnas de ceros y
     unos. Compara casos, reingresos, tasas, chi-cuadrado, V, valor p (con la distribución normal),
     dirección, decisión y descripción (leída con openpyxl).
  4. Tabla completa y lista. En todos los códigos: la tasa sin el código cuadra con los totales, la V
     con el chi-cuadrado y la decisión con los umbrales; la lista tiene exactamente los que pasan.

Solo lee el conjunto de desarrollo; con las rutas por defecto deja reportes/pruebas_filtro_codigos.md.

Uso (desde la carpeta Codigo):
    .venv\\Scripts\\python.exe 3.1_filtro_dimensionalidad\\pruebas\\verificar_14_filtro_codigos.py

Termina con código 0 si todas las comprobaciones se cumplen y con 1 si alguna falla.
"""

import csv
import hashlib
import math
import sys
import time
from datetime import date
from pathlib import Path
from statistics import NormalDist

import numpy as np
import pandas as pd
import pyarrow.compute as pc
import pyarrow.parquet as pq
from openpyxl import load_workbook

ACTIVIDAD = Path(__file__).resolve().parents[1]  # carpeta de la actividad
RAIZ = ACTIVIDAD.parent                          # carpeta Codigo
DESARROLLO = RAIZ / "data" / "processed" / "desarrollo.parquet"
RESULTADO = RAIZ / "data" / "processed" / "filtro_codigos.parquet"
REGISTRO = RAIZ / "docs" / "version_particion.csv"
DIR_TABLAS = RAIZ.parent / "Datos" / "Tablas auxiliares"
PARAMETROS = ACTIVIDAD / "reglas" / "parametros.csv"
LISTA = ACTIVIDAD / "reportes" / "14_codigos_retenidos.csv"
SALIDA = ACTIVIDAD / "reportes" / "pruebas_filtro_codigos.md"
COLUMNAS = {"DIAGNOSTICO": [f"DIAGNOSTICO{i}" for i in range(1, 36)],
            "PROCEDIMIENTO": [f"PROCEDIMIENTO{i}" for i in range(1, 31)]}
TABLAS = {"DIAGNOSTICO": "CIE-10.xlsx", "PROCEDIMIENTO": "CIE-9.xlsx"}
COLUMNAS_RESULTADO = ["VARIABLE", "BLOQUE", "CODIGO", "DESCRIPCION", "CASOS", "REINGRESOS", "TASA_CON_CODIGO",
                      "TASA_SIN_CODIGO", "CHI2", "VALOR_P", "V", "DIRECCION", "PASA"]
EJEMPLOS = ["N18.5", "K80.2", "U07.1", "Z37.0", "D61.3"]
TAMANO_MUESTRA = 300
SEMILLA = 2026
resultados = []


def comprobar(nombre: str, cumple, detalle: str = "") -> None:
    resultados.append((f"{nombre} {detalle}".rstrip(), bool(cumple)))
    print(f"[{'OK' if cumple else 'FALLA'}] {nombre} {detalle}".rstrip(), flush=True)


def leer_csv(ruta: Path) -> list[dict]:
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def sha256(ruta: Path) -> str:
    resumen = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            resumen.update(bloque)
    return resumen.hexdigest()


def fmt(n) -> str:
    return f"{int(n):,}".replace(",", ".")


def dec(x, decimales: int = 6) -> str:
    return f"{x:.{decimales}f}".replace(".", ",")


def iguales(a, b, rtol: float = 1e-9, atol: float = 1e-12) -> bool:
    """Compara dos listas de números con una tolerancia mínima (los decimales del computador no son exactos)."""
    return bool(np.allclose(np.asarray(a, dtype=float), np.asarray(b, dtype=float), rtol=rtol, atol=atol, equal_nan=True))


def main() -> None:
    inicio = time.time()
    par = {fila["parametro"]: fila["valor"].strip() for fila in leer_csv(PARAMETROS)}
    variables = [v for v in par["variables_objetivo"].split("|") if v]
    v_min = float(par["v_minima"].replace(",", "."))
    n_min = int(par["casos_minimos"])

    # ---------------------------------------------------------------- 1. versión y archivos
    registro = leer_csv(REGISTRO)[-1]
    comprobar(f"El conjunto de desarrollo tiene el hash de la versión {registro['version']} registrada de la partición",
              sha256(DESARROLLO) == registro["sha256_desarrollo"])
    todo = pd.read_parquet(RESULTADO)
    lista = pd.read_csv(LISTA, sep=";", decimal=",", encoding="utf-8-sig", dtype={"CODIGO": str})
    comprobar("filtro_codigos.parquet tiene las columnas esperadas, sin identificadores de egresos ni de pacientes",
              list(todo.columns) == COLUMNAS_RESULTADO)
    comprobar("La lista tiene las mismas columnas, salvo CHI2, VALOR_P y PASA",
              list(lista.columns) == [c for c in COLUMNAS_RESULTADO if c not in ("CHI2", "VALOR_P", "PASA")])

    # ---------------------------------------------------------------- 2. totales contados directamente
    t = pq.read_table(DESARROLLO, columns=[*variables, *COLUMNAS["DIAGNOSTICO"], *COLUMNAS["PROCEDIMIENTO"]],
                      filters=[("ES_CASO_ESTUDIO", "==", True)])
    comprobar("Se leen tantos casos de estudio como los registrados en la partición",
              t.num_rows == int(registro["casos_desarrollo"]), f"({fmt(t.num_rows)} casos)")
    con_dato = {k: pc.is_valid(t[k]).to_numpy(zero_copy_only=False) for k in variables}
    reingresa = {k: pc.fill_null(pc.equal(t[k], 1), False).to_numpy(zero_copy_only=False) for k in variables}
    N = {k: int(con_dato[k].sum()) for k in variables}    # casos con dato en cada variable
    R = {k: int(reingresa[k].sum()) for k in variables}   # reingresos de cada variable
    distintos = {b: set().union(*(pc.unique(t[c]).drop_null().to_pylist() for c in columnas))
                 for b, columnas in COLUMNAS.items()}
    filas_ok = all(((todo["VARIABLE"] == k) & (todo["BLOQUE"] == b)).sum() == len(distintos[b])
                   and set(todo.loc[(todo["VARIABLE"] == k) & (todo["BLOQUE"] == b), "CODIGO"]) == distintos[b]
                   for k in variables for b in COLUMNAS)
    comprobar("Cada variable tiene una fila por código distinto de cada bloque, contados columna por columna", filas_ok,
              f"({fmt(len(distintos['DIAGNOSTICO']))} diagnósticos y {fmt(len(distintos['PROCEDIMIENTO']))} procedimientos)")

    # ---------------------------------------------------------------- 3. muestra de códigos, por otro camino
    rng = np.random.default_rng(SEMILLA)
    todos = [(b, c) for b in COLUMNAS for c in sorted(distintos[b])]
    muestra = {todos[i] for i in rng.choice(len(todos), TAMANO_MUESTRA, replace=False)}
    muestra |= {("DIAGNOSTICO", c) for c in EJEMPLOS}
    for k in variables:  # los que quedaron justo en el límite de alguno de los dos umbrales, a cada lado
        sub = todo[todo["VARIABLE"] == k]
        con_n, con_v = sub[sub["CASOS"] >= n_min], sub[sub["V"] >= v_min]
        for parte in (con_n[con_n["V"] >= v_min].nsmallest(5, "V"), con_n[con_n["V"] < v_min].nlargest(5, "V"),
                      con_v[con_v["CASOS"] >= n_min].nsmallest(5, "CASOS"), con_v[con_v["CASOS"] < n_min].nlargest(5, "CASOS")):
            muestra |= set(zip(parte["BLOQUE"], parte["CODIGO"]))
    print(f"Recalculando {len(muestra)} códigos...", flush=True)

    # De cada columna se guardan solo las filas que tienen un código, y el código: así cada comparación
    # recorre menos datos (la columna DIAGNOSTICO30, por ejemplo, está casi siempre vacía).
    no_vacias = {b: [(np.flatnonzero(pc.is_valid(t[c]).to_numpy(zero_copy_only=False)), t[c].drop_null()) for c in columnas]
                 for b, columnas in COLUMNAS.items()}
    filas = []
    for i, (b, codigo) in enumerate(sorted(muestra), start=1):
        tiene = np.zeros(t.num_rows, dtype=bool)  # casos que tienen el código en alguna de sus columnas
        for filas_columna, codigos_columna in no_vacias[b]:
            tiene[filas_columna[pc.equal(codigos_columna, codigo).to_numpy(zero_copy_only=False)]] = True
        if i % 50 == 0:
            print(f"  {i} de {len(muestra)} códigos ({time.time() - inicio:.0f} s)", flush=True)
        for k in variables:
            x = tiene[con_dato[k]].astype(float)          # 1 si el caso tiene el código
            y = reingresa[k][con_dato[k]].astype(float)   # 1 si el caso reingresa
            n1, a = int(x.sum()), int(x @ y)
            v = abs(np.corrcoef(x, y)[0, 1]) if 0 < n1 < len(x) else 0.0  # en una tabla de 2×2, V = |correlación|
            filas.append({"VARIABLE": k, "BLOQUE": b, "CODIGO": codigo, "CASOS": n1, "REINGRESOS": a, "V": v,
                          "TASA_CON_CODIGO": 100 * a / n1 if n1 else np.nan,
                          "TASA_SIN_CODIGO": 100 * (R[k] - a) / (N[k] - n1), "CHI2": N[k] * v * v})
    calc = pd.DataFrame(filas)
    ref = todo.set_index(["VARIABLE", "BLOQUE", "CODIGO"]).loc[pd.MultiIndex.from_frame(calc[["VARIABLE", "BLOQUE", "CODIGO"]])]

    detalle = f"({len(muestra)} códigos × {len(variables)} variables)"
    comprobar("Muestra: casos y reingresos iguales a los del script",
              (ref["CASOS"].to_numpy() == calc["CASOS"].to_numpy()).all()
              and (ref["REINGRESOS"].to_numpy() == calc["REINGRESOS"].to_numpy()).all(), detalle)
    comprobar("Muestra: tasas de reingreso con el código y sin él iguales",
              iguales(ref["TASA_CON_CODIGO"], calc["TASA_CON_CODIGO"]) and iguales(ref["TASA_SIN_CODIGO"], calc["TASA_SIN_CODIGO"]))
    comprobar("Muestra: V igual a la correlación entre «tiene el código» y «reingresa»", iguales(ref["V"], calc["V"], rtol=1e-8))
    comprobar("Muestra: chi-cuadrado igual a N · V²", iguales(ref["CHI2"], calc["CHI2"], rtol=1e-8, atol=1e-9))
    # Un chi-cuadrado con 1 grado de libertad es el cuadrado de una normal estándar: p = 2 · (1 − Φ(raíz(chi²))).
    # Con chi-cuadrados grandes el valor p es tan chico que la resta pierde precisión; se comparan los menores que 40.
    chicos = calc["CHI2"].to_numpy() < 40
    p_normal = [2 * (1 - NormalDist().cdf(math.sqrt(x))) for x in calc["CHI2"].to_numpy()[chicos]]
    comprobar("Muestra: valor p igual al calculado con la distribución normal (chi-cuadrado menor que 40)",
              iguales(ref["VALOR_P"].to_numpy()[chicos], p_normal, rtol=1e-6), f"({int(chicos.sum())} filas)")
    direccion = np.select([calc["CASOS"] == 0, calc["TASA_CON_CODIGO"] > calc["TASA_SIN_CODIGO"],
                           calc["TASA_CON_CODIGO"] < calc["TASA_SIN_CODIGO"]], ["sin casos", "sube", "baja"], default="igual")
    pasa = ((calc["CASOS"] >= n_min) & (calc["V"] >= v_min)).to_numpy()
    comprobar("Muestra: dirección (sube o baja) y decisión (pasa o no) iguales",
              (ref["DIRECCION"].to_numpy() == direccion).all() and (ref["PASA"].to_numpy() == pasa).all(),
              f"({int(pasa.sum())} filas pasan)")
    descripciones = {}
    for b, archivo in TABLAS.items():
        libro = load_workbook(DIR_TABLAS / archivo, read_only=True)
        filas_xl = libro.active.iter_rows(values_only=True)
        encabezado = list(next(filas_xl))
        i_cod, i_desc = encabezado.index("Código"), encabezado.index("Descripción")
        descripciones[b] = {str(f[i_cod]).strip(): " ".join(str(f[i_desc] or "").split()) for f in filas_xl if f[i_cod] is not None}
        libro.close()
    esperada = np.array([descripciones[b].get(c, "") for b, c in zip(calc["BLOQUE"], calc["CODIGO"])])
    comprobar("Muestra: descripciones iguales a las de las tablas del DEIS (leídas con openpyxl)",
              (ref["DESCRIPCION"].to_numpy() == esperada).all())

    # ---------------------------------------------------------------- 4. tabla completa y lista
    cuadra, v_ok = True, True
    for k in variables:
        sub = todo[todo["VARIABLE"] == k]
        cuadra &= iguales(sub["TASA_SIN_CODIGO"], 100 * (R[k] - sub["REINGRESOS"]) / (N[k] - sub["CASOS"]))
        v_ok &= iguales(sub["V"], np.sqrt(sub["CHI2"] / N[k]), rtol=1e-12) and bool(sub["V"].between(0, 1).all())
    comprobar("Todos los códigos: la tasa sin el código cuadra con los casos con dato y los reingresos de cada variable", cuadra,
              "(" + "; ".join(f"{k}: {fmt(N[k])} casos con dato y {fmt(R[k])} reingresos" for k in variables) + ")")
    comprobar("Todos los códigos: la V es raíz(chi-cuadrado / N) y está entre 0 y 1", v_ok)
    pasa_todo = (todo["CASOS"] >= n_min) & (todo["V"] >= v_min)
    comprobar("Todos los códigos: pasan exactamente los que cumplen los dos umbrales",
              bool((todo["PASA"] == pasa_todo).all()), f"({int(pasa_todo.sum())} filas)")
    clave = ["VARIABLE", "BLOQUE", "CODIGO", "CASOS", "REINGRESOS"]
    de_lista = set(map(tuple, lista[clave].astype(str).to_numpy()))
    de_tabla = set(map(tuple, todo.loc[pasa_todo, clave].astype(str).to_numpy()))
    por_variable = lista["VARIABLE"].value_counts().reindex(variables, fill_value=0)
    comprobar("La lista tiene exactamente los códigos que pasan, con los mismos casos y reingresos",
              len(lista) == len(de_lista) and de_lista == de_tabla,
              f"({fmt(len(lista))} filas: " + ", ".join(f"{k} {fmt(n)}" for k, n in por_variable.items()) + ")")
    procedimientos = lista.loc[lista["BLOQUE"] == "PROCEDIMIENTO", "CODIGO"]
    comprobar("En la lista, los códigos de procedimiento conservan el cero inicial (dos dígitos antes del punto)",
              bool(procedimientos.str.fullmatch(r"\d\d(\.\d{1,2})?").all()))

    # ---------------------------------------------------------------- reporte
    todo_bien = all(ok for _, ok in resultados)
    print(f"\nRESULTADO: {'todas las comprobaciones se cumplen' if todo_bien else 'HAY FALLAS'} ({time.time() - inicio:.0f} s)")
    k_ej = "REING_30D" if "REING_30D" in variables else variables[0]
    ej = calc[(calc["VARIABLE"] == k_ej) & calc["CODIGO"].isin(EJEMPLOS)].set_index("CODIGO").reindex(EJEMPLOS)
    ej_ref = ref.reset_index().query("VARIABLE == @k_ej").set_index("CODIGO").reindex(EJEMPLOS)
    filas_ej = [f"| `{c}` | {fmt(ej_ref.loc[c, 'CASOS'])} / {fmt(ej.loc[c, 'CASOS'])} | "
                f"{fmt(ej_ref.loc[c, 'REINGRESOS'])} / {fmt(ej.loc[c, 'REINGRESOS'])} | "
                f"{dec(ej_ref.loc[c, 'V'])} / {dec(ej.loc[c, 'V'])} |" for c in EJEMPLOS]
    ruta = ACTIVIDAD.name
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text("\n\n".join([
        "# Prueba independiente del filtro de dimensionalidad",
        f"Generado por `{ruta}/pruebas/verificar_14_filtro_codigos.py` el {date.today():%d-%m-%Y}. El programa vuelve a calcular, "
        "sin usar ninguna función del script 14, lo que el script dejó en `data/processed/filtro_codigos.parquet` y en "
        f"`{ruta}/reportes/14_codigos_retenidos.csv`. Para una muestra de códigos ({len(muestra)}: {TAMANO_MUESTRA} al azar con la "
        f"semilla {SEMILLA}, los que quedaron justo en el límite de los umbrales y cinco ejemplos conocidos) marca los casos que "
        "tienen el código revisando columna por columna y calcula la V como la correlación entre dos columnas de ceros y unos. "
        "Solo lee el conjunto de desarrollo.",
        f"Partición probada: versión {registro['version']}, registrada el {registro['fecha']} en `docs/{REGISTRO.name}`.",
        f"**Resultado: {'todas las comprobaciones se cumplen' if todo_bien else 'HAY FALLAS'}** ({len(resultados)} comprobaciones, "
        f"{sum(not ok for _, ok in resultados)} fallas).",
        "\n".join(f"- {texto}: **{'sí' if ok else 'NO'}**." for texto, ok in resultados),
        f"Los cinco ejemplos con `{k_ej}` (script / recálculo):",
        "\n".join(["| Código | Casos | Reingresos | V |", "|---|---|---|---|", *filas_ej]),
    ]) + "\n", encoding="utf-8")
    print(f"Reporte guardado en {SALIDA}")
    sys.exit(0 if todo_bien else 1)


if __name__ == "__main__":
    main()
